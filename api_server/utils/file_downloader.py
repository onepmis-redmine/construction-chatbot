import os
from pathlib import Path
from fastapi import HTTPException, Response
from fastapi.responses import FileResponse
from utils.logger import get_logger
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import tempfile
import time
from datetime import datetime

logger = get_logger(__name__)

class FileDownloader:
    def __init__(self, base_dir: Path):
        self.base_dir = base_dir
        self.questions_dir = base_dir / "questions"
        self.docs_dir = base_dir / "docs"
    
    def download_questions(self) -> Response:
        """저장된 질문 데이터를 Excel 파일로 다운로드합니다."""
        temp_path = None
        try:
            questions_file = self.questions_dir / "saved_questions.csv"
            
            if not questions_file.exists():
                raise HTTPException(status_code=404, detail="저장된 질문 파일이 없습니다.")
            
            # CSV 파일을 DataFrame으로 읽기
            df = pd.read_csv(questions_file, encoding='utf-8')
            
            # 임시 파일 생성
            temp_fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
            os.close(temp_fd)  # 파일 디스크립터 닫기
            
            # DataFrame을 Excel로 저장
            df.to_excel(temp_path, index=False, engine='openpyxl')
            
            # Excel 파일 열어서 스타일 조정
            wb = load_workbook(temp_path)
            ws = wb.active
            
            # 열 너비 자동 조정
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                
                for cell in column:
                    if cell.value:
                        # 줄바꿈 문자를 기준으로 최대 길이 계산
                        lines = str(cell.value).split('\n')
                        max_line_length = max(len(line) for line in lines)
                        max_length = max(max_length, max_line_length)
                
                # 열 너비 설정 (최소 10, 최대 50)
                adjusted_width = min(max(max_length + 2, 10), 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # 헤더 스타일 설정
            header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            header_font = Font(bold=True)
            
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            # 데이터 셀 스타일 설정
            for row in ws.iter_rows(min_row=2):
                max_height = 0
                for cell in row:
                    if cell.value:
                        # 줄바꿈 문자 개수에 따라 높이 조정
                        lines = str(cell.value).split('\n')
                        line_count = len(lines)
                        cell_height = line_count * 25  # 기본 높이 25에 줄 수만큼 곱함
                        max_height = max(max_height, cell_height)
                        
                        # 셀 정렬 설정
                        cell.alignment = Alignment(
                            horizontal='left',
                            vertical='center',
                            wrap_text=True
                        )
                
                if max_height > 0:
                    ws.row_dimensions[row[0].row].height = max_height
            
            # 조정된 Excel 파일 저장
            wb.save(temp_path)
            
            # 파일이 완전히 저장될 때까지 잠시 대기
            time.sleep(0.1)
            
            # Excel 파일을 읽어서 응답으로 전송
            with open(temp_path, 'rb') as f:
                content = f.read()
            
            return Response(
                content=content,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": "attachment; filename=saved_questions.xlsx"
                }
            )
        except Exception as e:
            logger.error(f"질문 다운로드 중 오류 발생: {e}")
            raise HTTPException(status_code=500, detail=str(e))
        finally:
            # 임시 파일 정리
            if temp_path and os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception as e:
                    logger.error(f"임시 파일 삭제 중 오류 발생: {e}")

    def download_enhanced_qa(self) -> FileResponse:
        """구조화된 FAQ 파일을 다운로드합니다."""
        try:
            source_filename = "enhanced_qa_pairs.xlsx"
            file_path = self.docs_dir / source_filename
            if not file_path.exists():
                raise HTTPException(status_code=404, detail="파일을 찾을 수 없습니다.")
            
            # 다운로드할 파일명에 타임스탬프 추가
            timestamp = datetime.now().strftime("%Y%m%d%H%M")
            download_filename = f"enhanced_qa_pairs_{timestamp}.xlsx"
            
            return FileResponse(
                path=file_path,
                filename=download_filename,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            logger.error(f"파일 다운로드 중 오류 발생: {str(e)}")
            raise HTTPException(status_code=500, detail=str(e)) 